from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

doc.add_heading('THISO Leasing Platform — Parking Dashboard UI', level=0)

p = doc.add_paragraph()
p.add_run('Module: ').bold = True
p.add_run('Bãi đỗ xe (Parking) — Báo cáo bãi đỗ xe / Giao dịch bãi đỗ xe')
p2 = doc.add_paragraph()
p2.add_run('Ngày chụp: ').bold = True
p2.add_run('2026-08-25')
p3 = doc.add_paragraph()
p3.add_run('Nguồn dữ liệu: ').bold = True
p3.add_run('CSDL bãi đỗ xe thực (sync-parking, remote 125.234.136.82) qua SSH tunnel, không phải dữ liệu mẫu.')

doc.add_heading('1. Báo cáo bãi đỗ xe (Legend / Dashboard) — Bãi Phan Huy Ích', level=1)
doc.add_paragraph('Tổng quan doanh thu tháng này / tháng trước theo hình thức thanh toán (Tiền Mặt, Online, Voucher), '
                   'số lượt xe, và biểu đồ doanh thu / lượt xe theo năm (khoảng 2024–2026).')
doc.add_picture('docs/parking-dashboard/parking-report-dashboard.png', width=Inches(6.3))

doc.add_heading('2. Báo cáo bãi đỗ xe — Bãi Phan Văn Trị', level=1)
doc.add_paragraph('Cùng màn hình báo cáo, chuyển sang bãi đỗ "Phan Văn Trị" — bãi có dữ liệu đồng bộ mới nhất '
                   '(2026-02-23 đến 2026-08-01), gần với thời điểm hiện tại nhất trong 3 bãi.')
doc.add_picture('docs/parking-dashboard/parking-report-dashboard-2.png', width=Inches(6.3))

doc.add_heading('3. Giao dịch bãi đỗ xe (Transactions) — Bãi Phan Văn Trị', level=1)
doc.add_paragraph('Danh sách giao dịch chi tiết, lọc theo khoảng ngày 2026-07-15 → 2026-07-16, khoảng thời gian '
                   'có dữ liệu giao dịch thực tế (bãi Phan Văn Trị có ~1.37 triệu bản ghi trong toàn bộ dải ngày).')
doc.add_picture('docs/parking-dashboard/parking-transaction-list.png', width=Inches(6.3))

doc.add_heading('4. Ghi chú kỹ thuật', level=1)
notes = [
    'Kết nối CSDL bãi đỗ qua SSH tunnel tới root@125.234.136.82, forward local port 5433 → remote port 5433 '
    '(khớp PARKING_DATABASE_URL trong .env gốc: host.docker.internal:5433).',
    'Dữ liệu 3 bãi đỗ (tenants trong CSDL sync-parking):',
    '  • Phan Huy Ích: 2023-12-04 → 2024-07-17 (654,908 giao dịch)',
    '  • Sala: 2024-07-20 → 2024-09-28 (349,594 giao dịch)',
    '  • Phan Văn Trị: 2026-02-23 → 2026-08-01 (1,372,233 giao dịch) — dữ liệu mới nhất',
    'Backend chạy Docker dev mode (leasing-backend, port 3050), frontend Vite dev server (leasing-frontend, port 5173).',
]
for n in notes:
    doc.add_paragraph(n, style='List Bullet' if not n.startswith('  ') else None)

doc.save('docs/parking-dashboard/Parking-Dashboard-UI.docx')
print('saved')
